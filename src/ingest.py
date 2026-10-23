import argparse
import pickle
import re
import sqlite3
from collections import Counter
from pathlib import Path

import faiss
from rank_bm25 import BM25Okapi

import config
from embedder import Embedder
from topics import detect_topics, tokenize


def parse_front_matter(text):
    meta = {}
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
        text = text[m.end():]
    return meta, text


def load_markdown(path: Path):
    meta, text = parse_front_matter(path.read_text(encoding="utf-8", errors="ignore"))
    pages, section, buf = [], meta.get("title", path.stem), []
    for line in text.splitlines():
        h = re.match(r"^(#{1,3})\s+(.*)", line)
        if h:
            if buf:
                pages.append((None, section, "\n".join(buf)))
            section, buf = h.group(2).strip(), []
        else:
            buf.append(line)
    if buf:
        pages.append((None, section, "\n".join(buf)))
    return meta, pages


SECTION_RE = re.compile(r"^\s*(\d+(?:\.\d+)+)\s+([A-Z][^\n]{2,70})$")


def _looks_like_heading(title):
    words = [w for w in re.findall(r"[A-Za-z][\w'-]*", title) if len(w) > 3]
    if not words or len(title) > 60 or title.rstrip().endswith((".", ",", "?")):
        return False
    return sum(w[0].isupper() for w in words) / len(words) >= 0.5


def load_pdf(path: Path, keep=None):
    import pdfplumber
    pages, section = [], None
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            if keep is not None and i not in keep:
                continue
            txt = page.extract_text(x_tolerance=1.5) or ""
            if len(txt.strip()) < 20:
                txt = ocr_page(page) or txt
            buf, start_section = [], section
            for line in txt.splitlines():
                m = SECTION_RE.match(line)
                if m and not _looks_like_heading(m.group(2)):
                    m = None
                if m and buf:
                    pages.append((i, start_section, "\n".join(buf)))
                    buf, start_section = [], None
                if m:
                    section = f"{m.group(1)} {m.group(2).strip()}"
                    start_section = section
                    continue
                buf.append(line)
            if buf:
                pages.append((i, start_section, "\n".join(buf)))
    return {}, strip_repeated_lines(pages)


def ocr_page(page):
    try:
        import pytesseract
        img = page.to_image(resolution=300).original
        return pytesseract.image_to_string(img)
    except Exception:
        return ""


_OK_CHARS = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789 (){};&*=<>+-_,.\"'#/%![]:|\\?")


def _plausible(line):
    t = line.strip()
    if len(t) < 3:
        return t in ("{", "}", "};")
    ok = sum(c in _OK_CHARS for c in t) / len(t)
    word = re.search(r"[A-Za-z_]*[aeiouAEIOU_][A-Za-z_]*", t)
    return ok >= 0.9 and bool(word and len(word.group(0)) >= 3)


def _ocr_blob(blob):
    try:
        import io
        import pytesseract
        from PIL import Image
        img = Image.open(io.BytesIO(blob))
        if img.width < 200 or img.height < 80:
            return ""
        g = img.convert("L")
        g = g.resize((g.width * 2, g.height * 2))
        txt = pytesseract.image_to_string(g, config="--psm 4")
        return "\n".join(l for l in txt.splitlines() if _plausible(l))
    except Exception:
        return ""


def _shape_texts(shapes, part):
    texts, images = [], []
    for shp in shapes:
        if shp.shape_type == 6 and hasattr(shp, "shapes"):
            t, im = _shape_texts(shp.shapes, part)
            texts += t
            images += im
            continue
        if getattr(shp, "has_table", False) and shp.has_table:
            for row in shp.table.rows:
                cells = [c.text.strip() for c in row.cells]
                if sum(1 for c in cells if c) >= 2:
                    texts.append(" | ".join(cells))
            continue
        if shp.has_text_frame and shp.text_frame.text.strip():
            texts.append(shp.text_frame.text)
        for rid in re.findall(r'r:embed="(\w+)"', shp._element.xml):
            try:
                images.append(part.related_part(rid).blob)
            except Exception:
                pass
    return texts, images


def load_pptx(path: Path):
    from pptx import Presentation
    pages = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else None
        texts, images = _shape_texts(slide.shapes, slide.part)
        body = [t for t in texts if t.strip() != (title or "")]
        if sum(len(t.split()) for t in body) < 25 and images:
            ocr = "\n".join(o for o in (_ocr_blob(b) for b in images[:3]) if o.strip())
            if len(ocr.split()) >= 5:
                body.append("[Text read from slide image]\n" + ocr)
        if slide.has_notes_slide:
            body.append(slide.notes_slide.notes_text_frame.text)
        pages.append((i, title, "\n".join(body)))
    return {}, pages


def load_ppt(path: Path):
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("install LibreOffice to read .ppt files (or save them as .pptx)")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, "--headless", "--convert-to", "pptx", str(path), "--outdir", tmp],
                       check=True, capture_output=True, timeout=180)
        return load_pptx(Path(tmp) / (path.stem + ".pptx"))


def load_docx(path: Path):
    import docx
    d = docx.Document(str(path))
    pages, section, buf = [], None, []
    for para in d.paragraphs:
        if para.style.name.lower().startswith("heading") and para.text.strip():
            if buf:
                pages.append((None, section, "\n".join(buf)))
            section, buf = para.text.strip(), []
        else:
            buf.append(para.text)
    for t in d.tables:
        buf += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    if buf:
        pages.append((None, section, "\n".join(buf)))
    return {}, pages


def load_txt(path: Path):
    meta, text = parse_front_matter(path.read_text(encoding="utf-8", errors="ignore"))
    return meta, [(None, path.stem, text)]


LOADERS = {".md": load_markdown, ".pdf": load_pdf, ".pptx": load_pptx, ".ppt": load_ppt,
           ".docx": load_docx, ".txt": load_txt}


def strip_repeated_lines(pages):
    n_pages = len({p for p, _, _ in pages})
    if n_pages < 4:
        return pages
    norm = lambda l: re.sub(r"\d+", "#", l.strip())
    per_page = {}
    for p, _, t in pages:
        per_page.setdefault(p, set()).update(norm(l) for l in t.splitlines() if 0 < len(l.strip()) < 80)
    counts = Counter(l for lines in per_page.values() for l in lines)
    repeated = {l for l, c in counts.items() if c >= max(3, 0.3 * n_pages)}
    return [(p, s, "\n".join(l for l in t.splitlines() if norm(l) not in repeated))
            for p, s, t in pages]


def clean(text: str) -> str:
    text = re.sub(r"<number>", "", text)
    text = "\n".join(l for l in text.splitlines()
                     if not re.search(r"[\uac00-\ud7af]", l))
    text = re.sub(r"\((\w)\)\s*\n\s*([a-z])\s*(?=\n|$)", r"(\1_\2)", text)
    text = re.sub(r"(?m)^\s*(page\s*)?\d+\s*(of\s*\d+)?\s*$", "", text, flags=re.I)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()


FILLER = re.compile(r"subscribe|bell button|like the video|like button|share it|share this|welcome to gate smashers|"
                    r"hello (friends|students|guys)|thank you|description box|comment box|playlist|"
                    r"latest (video|notification)", re.I)


def clean_transcript(text):
    cues = re.split(r"(?m)^\s*\d{1,2}:\d{2}(?::\d{2})?\s*", text)
    kept = [" ".join(c.split()) for c in cues if c.strip() and not FILLER.search(c)]
    words = " ".join(kept).split()
    return "\n".join(" ".join(words[i:i + 60]) for i in range(0, len(words), 60))


def load_sources_meta(directory: Path):
    f = directory / "_sources.json"
    if f.exists():
        import json
        return json.loads(f.read_text())
    return {}


def shingles(text, n=5):
    w = re.findall(r"[a-z0-9]+", text.lower())
    return {" ".join(w[i:i + n]) for i in range(max(0, len(w) - n + 1))}


def chunk_words(text, size=config.CHUNK_WORDS, overlap=config.CHUNK_OVERLAP):
    lines = []
    for line in text.splitlines():
        w = line.split()
        while len(w) > size:
            lines.append(" ".join(w[:size]))
            w = w[size - overlap:]
        if w:
            lines.append(" ".join(w))
    if sum(len(l.split()) for l in lines) <= size:
        return ["\n".join(lines)] if lines else []
    out, cur, n = [], [], 0
    for line in lines:
        lw = len(line.split())
        if cur and n + lw > size:
            out.append("\n".join(cur))
            keep, kn = [], 0
            for prev in reversed(cur):
                kn += len(prev.split())
                if kn > overlap:
                    break
                keep.insert(0, prev)
            cur, n = keep, sum(len(x.split()) for x in keep)
        cur.append(line)
        n += lw
    if cur:
        out.append("\n".join(cur))
    return out


def label_topic(meta_topic, section, text, file_topic=None, doc_type=None):
    if meta_topic and meta_topic != "auto":
        return meta_topic
    if doc_type == "question_bank":
        found = detect_topics(text, top_n=1)
        if found:
            return found[0]
    for t in config.TOPICS:
        if section and t.lower() in section.lower():
            return t
    found = detect_topics(f"{section or ''} {text}", top_n=1)
    return found[0] if found else (file_topic or "General")


CACHE_DIR = config.INDEX_DIR / "cache"


def load_cached(path, ranges=None):
    import hashlib
    import pickle as pk
    st = path.stat()
    key = hashlib.md5(f"{path.resolve()}|{st.st_size}|{st.st_mtime}|{ranges}|v2".encode()).hexdigest()
    f = CACHE_DIR / f"{key}.pkl"
    if f.exists():
        with open(f, "rb") as fh:
            return pk.load(fh)
    if ranges and path.suffix.lower() == ".pdf":
        keep = {p for a, b, _ in ranges for p in range(a, b + 1)}
        out = load_pdf(path, keep=keep)
    else:
        out = LOADERS[path.suffix.lower()](path)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    with open(f, "wb") as fh:
        pk.dump(out, fh)
    return out


def build_chunks():
    files = []
    for d, default_type in [(config.NOTES_DIR, "notes"), (config.QUESTIONS_DIR, "question_bank"),
                            (config.RAW_DIR, "material")]:
        for p in sorted(d.rglob("*")):
            if p.suffix.lower() in LOADERS:
                files.append((p, default_type))

    chunks, seen_docs = [], []
    side = {}
    for path, _ in files:
        if path.parent not in side:
            side[path.parent] = load_sources_meta(path.parent)
    for path, default_type in files:
        try:
            side_meta = side.get(path.parent, {}).get(path.name, {})
            if side_meta.get("skip"):
                continue
            ranges = side_meta.get("pages")
            meta, pages = load_cached(path, ranges)
            if side_meta.get("transcript"):
                pages = [(pg, side_meta.get("title", sec), clean_transcript(t)) for pg, sec, t in pages]
        except Exception as e:
            print(f"  ! skipped {path.name}: {e}")
            continue
        meta = {**side.get(path.parent, {}).get(path.name, {}), **meta}
        if meta.get("skip"):
            continue
        sh = shingles(" ".join(t for _, _, t in pages))
        dup = next((n for n, other in seen_docs if sh and len(sh & other) / max(1, max(len(sh), len(other))) > 0.6), None)
        if dup:
            print(f"  = {path.name:<32} skipped (same content as {dup})")
            continue
        seen_docs.append((path.name, sh))
        doc_type = meta.get("doc_type", default_type)
        source = meta.get("source", path.stem.replace("_", " "))
        file_topic = meta.get("fallback_topic")
        if not file_topic and (not meta.get("topic") or meta.get("topic") == "auto"):
            whole = " ".join(t for _, _, t in pages)[:20000]
            ft = detect_topics(whole, top_n=1)
            file_topic = ft[0] if ft else None

        n_before = len(chunks)
        merged, buf = [], None
        for page, section, text in pages:
            text = clean(text)
            if not text:
                continue
            bw, tw = (len(buf[2].split()) if buf else 0), len(text.split())
            if buf and buf[1] == section and bw + tw <= config.CHUNK_WORDS:
                buf = (buf[0], section, buf[2] + "\n" + text)
            elif buf and bw < 40 and doc_type != "question_bank":
                sec = buf[1] if not section or section == buf[1] else (f"{buf[1]} / {section}" if buf[1] else section)
                buf = (buf[0], sec, buf[2] + "\n" + text)
            else:
                if buf:
                    merged.append(buf)
                buf = (page, section, text)
        if buf:
            merged.append(buf)

        for page, section, text in merged:
            if doc_type == "question_bank":
                pieces = [q.strip() for q in re.split(r"\n(?=Q[:.\d])", "\n" + text) if q.strip()]
            else:
                pieces = chunk_words(text)
            for piece in pieces:
                if ranges and page:
                    pt = next((t for a, b, t in ranges if a <= page <= b), "auto")
                    if pt.startswith("auto"):
                        topic = label_topic("auto", section, piece, pt.partition(":")[2] or file_topic, doc_type)
                    else:
                        topic = pt
                else:
                    topic = label_topic(meta.get("topic"), section, piece, file_topic, doc_type)
                chunks.append({
                    "text": piece, "topic": topic, "section": section or "",
                    "source": source, "file": path.name, "page": page if doc_type != "notes" else None,
                    "doc_type": doc_type,
                })
        print(f"  + {path.name:<32} {len(chunks) - n_before:>3} chunks  [{doc_type}]")
    return chunks


def index_text(c):
    return f"{c['topic']}. {c['section']}. {c['text']}"


def main(backend=None):
    config.INDEX_DIR.mkdir(exist_ok=True)
    print("Loading and chunking material...")
    chunks = build_chunks()
    if not chunks:
        raise SystemExit("No documents found. Put files in data/notes, data/questions or data/raw.")
    texts = [index_text(c) for c in chunks]

    if config.META_DB.exists():
        config.META_DB.unlink()
    con = sqlite3.connect(config.META_DB)
    con.execute("""CREATE TABLE chunks (id INTEGER PRIMARY KEY, text TEXT, topic TEXT,
                   section TEXT, source TEXT, file TEXT, page INTEGER, doc_type TEXT)""")
    con.executemany("INSERT INTO chunks VALUES (?,?,?,?,?,?,?,?)",
                    [(i, c["text"], c["topic"], c["section"], c["source"], c["file"],
                      c["page"], c["doc_type"]) for i, c in enumerate(chunks)])
    con.commit()
    con.close()

    with open(config.BM25_PATH, "wb") as f:
        pickle.dump(BM25Okapi([tokenize(t) for t in texts]), f)

    emb = Embedder(backend).fit(texts)
    vecs = emb.encode_docs(texts)
    index = faiss.IndexFlatIP(vecs.shape[1])
    index.add(vecs)
    faiss.write_index(index, str(config.FAISS_PATH))
    emb.save()

    by_topic = Counter(c["topic"] for c in chunks)
    print(f"\nIndexed {len(chunks)} chunks from {len({c['file'] for c in chunks})} files "
          f"| embeddings: {emb.name} (dim {vecs.shape[1]})")
    for t, n in by_topic.most_common():
        print(f"  {t:<30} {n}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--backend", choices=["auto", "st", "tfidf"], default=None)
    main(ap.parse_args().backend)
